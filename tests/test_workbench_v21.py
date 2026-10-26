"""User task, complete script, lineage and export regressions (no commercial calls)."""
import io
import json
import zipfile
from copy import deepcopy
from pathlib import Path

import pytest
from creative_agent.store import Store, Conflict
from creative_agent.scripts import Scripts


@pytest.fixture
def core(tmp_path):
    s=Store(tmp_path/'db'); p=s.create_project('商品项目')['id']; q=s.create_project('其他')['id']
    return s,Scripts(s),p,q


def full(title='生活场景脚本'):
    return dict(title=title,objective='用动作说明使用方式',rows=[dict(start_ms=0,end_ms=4500,dialogue='先把桌上的东西收好。',visual='人物将桌面杂物放进盒子',tone='自然',references=[],own_references=[],reshoot='补拍整理桌面的动作',claim_ids=[],subtitle='开始整理',camera='固定中景',props=['收纳盒'])],confirmed=False,usage_status='planned',feedback_reason='')


def test_save_full_scripts_preserves_timings_and_idempotency(core):
    s,svc,p,q=core
    session=s.put(p,'session',dict(scope={'project_id':p,'asset_ids':[]},turns=[dict(question='写两个脚本',answer={'result':{'scripts':[full('A'),full('B')]}},evidence=[])]))
    a=svc.from_session(p,session['id'],0,script_index=1)
    assert a['title']=='B'
    assert a['rows'][0]['end_ms']==4500
    assert a['rows'][0]['visual']=='人物将桌面杂物放进盒子'
    assert a['rows'][0]['reshoot']=='补拍整理桌面的动作'
    assert svc.from_session(p,session['id'],0,script_index=1)['id']==a['id']
    revised=svc.revise(p,a['id'],1,{**full('人工改稿'),'rows':[{**a['rows'][0],'dialogue':'人工修订'}]})
    assert svc.from_session(p,session['id'],0,script_index=1)['rows'][0]['dialogue']=='人工修订'
    assert revised['version']==2 and revised['source']==a['source']


def test_outline_is_not_silently_converted_to_full_shooting_script(core):
    s,svc,p,q=core
    session=s.put(p,'session',{'scope':{'asset_ids':[]},'turns':[{'answer':{'result':{'script_outline':['开头','展示']}},'evidence':[]}]})
    with pytest.raises(ValueError,match='完整脚本'):svc.from_session(p,session['id'],0)
    assert s.list(p,'script')==[]


def test_five_exports_have_full_content_and_exact_revision(core):
    s,svc,p,q=core
    script=svc.create(p,full('中文拍摄脚本'))
    svc.revise(p,script['id'],1,full('新标题'))
    for fmt in ['md','json','docx','xlsx','pdf']:
        data=svc.export(p,script['id'],fmt,version=1)
        assert isinstance(data,bytes) and data
        if fmt=='json':assert json.loads(data)['title']=='中文拍摄脚本'
        if fmt=='md':assert '补拍整理桌面的动作' in data.decode()
        if fmt=='docx':
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                xml=z.read('word/document.xml').decode()
                assert '中文拍摄脚本' in xml and '先把桌上的东西收好。' in xml and '补拍整理桌面的动作' in xml
        if fmt=='xlsx':
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data)); text=str(list(wb.active.values))
            assert '先把桌上的东西收好。' in text and '固定中景' in text
        if fmt=='pdf':assert data.startswith(b'%PDF')


def test_human_readable_exports_include_actual_product_evidence(core):
    s,svc,p,q=core
    svc.save_facts(p,0,{'product':'收纳盒','claims':[{'fact_id':'color','statement':'白色外壳','source':'包装标识2026版','confirmed':True}]})
    body=full();body['rows'][0]['factual_claims']=[{'statement':'白色外壳','fact_ids':['color'],'reviewed':True}]
    script=svc.create(p,body)
    for fmt in ['md','docx','xlsx']:
        data=svc.export(p,script['id'],fmt)
        if fmt=='md':text=data.decode()
        elif fmt=='docx':
            with zipfile.ZipFile(io.BytesIO(data)) as z:text=z.read('word/document.xml').decode()
        else:
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(data));text=str([list(sheet.values) for sheet in book])
        assert '白色外壳' in text and '包装标识2026版' in text and '依据ID：color' in text
    assert b'/FontFile2' in svc.export(p,script['id'],'pdf'), 'PDF must embed Chinese glyphs rather than require viewer language packs'


def test_legacy_field_template_remains_usable_without_rewriting_record(core,tmp_path):
    from creative_agent.media import MediaService
    from creative_agent.workflows import Workflows
    from PIL import Image
    s,svc,p,q=core;media=MediaService(s,tmp_path/'media');path=tmp_path/'image.png';Image.new('RGB',(20,20),'blue').save(path)
    asset=media.ingest(p,path)
    legacy=s.put(p,'template',{'title':'旧模板','fields':['主体']})
    class Model:
        def chat(self,*args,**kwargs):return {'data':{'summary':'蓝色图片','fields':{'主体':'蓝色区域'},'scenes':[]},'usage':{}}
    parsed=Workflows(s,media).parse(p,[asset['id']],['主体'],Model(),template_snapshot=legacy)
    assert list(parsed['rows'][0]['cells'].values())[0]['value']=='蓝色区域'
    assert 'definitions' not in s.get(p,'template',legacy['id'])


def test_own_references_reject_reference_material_wrong_sku_and_time(core,tmp_path):
    s,svc,p,q=core
    path=tmp_path/'local.mp4';path.write_bytes(b'local media locator only')
    a=s.put(p,'asset',{'source_kind':'own','file_path':str(path),'inspection':'ready','media_type':'video','metadata':{'duration_ms':5000,'sku_id':'sku-B'}})
    body={**full(),'sku_id':'sku-A'}
    body['rows'][0]['own_references']=[{'asset_id':a['id'],'start_ms':0,'end_ms':3000}]
    with pytest.raises(ValueError,match='SKU|款式'):svc.create(p,body)
    a['metadata']['sku_id']='sku-A';a['source_kind']='inspiration';s.put(p,'asset',a,a['id'])
    with pytest.raises(ValueError,match='自有|研究'):svc.create(p,body)
    a['source_kind']='own';s.put(p,'asset',a,a['id']);body['rows'][0]['own_references'][0]['end_ms']=6000
    with pytest.raises(ValueError,match='时间|范围'):svc.create(p,body)


def test_facts_support_is_not_proved_by_an_arbitrary_existing_id(core):
    s,svc,p,q=core
    facts=svc.save_facts(p,0,{'product':'收纳盒','sku_id':'A','claims':[{'fact_id':'color','statement':'白色','source':'包装','confirmed':True}]})
    body={**full(),'sku_id':'A','product_facts_snapshot':facts,'confirmed':True}
    body['rows'][0].update(dialogue='可以承重50公斤',claim_ids=['color'],factual_claims=[{'statement':'承重50公斤','fact_ids':['color']}])
    with pytest.raises(ValueError,match='依据|核实|事实|支持'):svc.create(p,body)
    body['confirmed']=False
    draft=svc.create(p,body)
    assert draft['unresolved_items']


def test_expired_fact_cannot_confirm_and_fact_snapshot_cannot_be_spoofed(core):
    s,svc,p,q=core
    facts=svc.save_facts(p,0,{'product':'食品','claims':[{'fact_id':'promo','statement':'今天买一送一','source':'促销单','confirmed':True,'valid_until':'2000-01-01'}]})
    body={**full(),'product_facts_snapshot':facts,'confirmed':True}
    body['rows'][0].update(dialogue='今天买一送一',claim_ids=['promo'])
    with pytest.raises(ValueError,match='过期|有效期|核实'):svc.create(p,body)
    body['product_facts_snapshot']['claims'][0]['valid_until']='2999-01-01'
    with pytest.raises(ValueError,match='快照|事实|版本'):svc.create(p,body)


def test_task_spec_does_not_require_three_concepts_or_one_script_ten_rows():
    from creative_agent.creative_tasks import TaskSpec, validate_delivery
    spec=TaskSpec(output_type='scripts',count=2,target_duration_ms=4500)
    result=validate_delivery({'scripts':[full()]},spec)
    assert result['delivery_progress']=={'requested':2,'completed':1,'state':'partial'}
    result=validate_delivery({'scripts':[full('A'),full('B')]},spec)
    assert result['delivery_progress']['state']=='complete'
    assert validate_delivery({'comparisons':['开头的展示顺序不同']},TaskSpec(output_type='comparison'))['delivery_progress']['completed']==1
    with pytest.raises(ValueError):validate_delivery({'scripts':[{'title':'只有标题'}]},spec)


def test_selection_freezes_own_search_manifest_and_rejects_cross_project(core):
    from creative_agent.creative_tasks import freeze_selection,TaskSpec
    s,svc,p,q=core
    a=s.put(p,'asset',{'source_kind':'own','metadata':{'duration_ms':2000},'sha256':'one'})
    b=s.put(q,'asset',{'source_kind':'own','metadata':{'duration_ms':2000}})
    snap=freeze_selection(s,p,{'project_id':p,'asset_ids':[]},TaskSpec(own_asset_ids=[a['id']]))
    assert snap['own_assets'][0]['id']==a['id'] and snap['own_assets'][0]['sha256']=='one'
    s.put(p,'asset',{'source_kind':'own','metadata':{},'sha256':'two'})
    assert len(snap['own_assets'])==1
    with pytest.raises(ValueError):freeze_selection(s,p,{},TaskSpec(own_asset_ids=[b['id']]))


def test_agent_saves_partial_typed_delivery_and_keeps_snapshot(core,tmp_path):
    from creative_agent.agent import AnalysisAgent
    from creative_agent.creative_tasks import TaskSpec
    from creative_agent.contracts import AnalysisScope
    s,svc,p,q=core
    class Model:
        def chat(self,system,question,**kwargs):
            return {'data':{'summary':'完成一个，尚有一个','findings':[],'result':{'scripts':[full()]}},'usage':{}}
    result=AnalysisAgent(s,None).analyze(AnalysisScope(project_id=p),'inspiration','写两个完整脚本',Model(),discussion_only=True,task_spec=TaskSpec(output_type='scripts',count=2))
    assert result['turns'][0]['answer']['result']['delivery_progress']['state']=='partial'
    assert result['turns'][0]['task_spec']['count']==2
    assert result['selection_snapshot']['project_id']==p


def test_local_preferences_and_full_script_routes_are_project_scoped(tmp_path):
    from fastapi.testclient import TestClient
    from creative_agent.api import create_app
    app=create_app(tmp_path,token='test',run_jobs=False)
    with TestClient(app,base_url='http://127.0.0.1:8765') as client:
        h={'Authorization':'Bearer test'}
        p=client.post('/api/projects',headers=h,json={'name':'p'}).json()['id']
        q=client.post('/api/projects',headers=h,json={'name':'q'}).json()['id']
        root='/api/projects/'+p
        r=client.put(root+'/preferences',headers=h,json={'default_quote':'2','default_binding_id':''})
        assert r.status_code==200,r.text
        assert client.get('/api/projects/'+q+'/preferences',headers=h).json()['default_quote']==''
        script=client.post(root+'/scripts',headers=h,json=full()).json()
        assert client.get(root+'/state',headers=h).json()['scripts'][0]['id']==script['id']
        doc=client.get(root+'/scripts/'+script['id']+'/export?format=docx&version=1',headers=h)
        assert doc.status_code==200 and doc.content.startswith(b'PK')
        assert client.get('/api/projects/'+q+'/scripts/'+script['id']+'/export?format=docx',headers=h).status_code==404


def test_template_is_not_a_prompt_and_revisions_preserve_stable_fields(core):
    from creative_agent.analysis_templates import Templates
    s,svc,p,q=core;t=Templates(s)
    first=t.save(p,{'title':'食品镜头','fields':['开头','场景']})
    assert not s.list(p,'instruction')
    second=t.save(p,{'title':'新版名称','fields':['场景','开头']},record_id=first['id'],expected_version=1)
    assert second['version']==2
    assert {d['name']:d['field_id'] for d in first['definitions']}=={d['name']:d['field_id'] for d in second['definitions']}
    assert second['revisions'][0]['title']=='食品镜头'


def test_own_footage_search_hard_filters_before_matching(core,tmp_path):
    from creative_agent.footage import FootageSearch
    from creative_agent.creative_tasks import freeze_selection,TaskSpec
    s,svc,p,q=core
    ids=[]
    for sku,kind in [('A','own'),('B','own'),('A','inspiration')]:
        f=tmp_path/(sku+kind);f.write_bytes(b'local')
        a=s.put(p,'asset',{'source_kind':kind,'media_type':'video','path':str(f),'inspection':'ready','sha256':__import__('hashlib').sha256(f.read_bytes()).hexdigest(),'metadata':{'sku_id':sku,'duration_ms':5000,'reuse_allowed':True,'transcript_source':'user_import','subtitles':[{'start_ms':0,'end_ms':2000,'text':'把收纳盒放在桌上'}]}})
        ids.append(a['id'])
    snap=freeze_selection(s,p,{},TaskSpec(own_asset_ids=ids))
    result=FootageSearch(s,None).search(p,'桌上收纳盒',snap,sku_id='A')
    assert len(result)==1 and result[0]['asset_id']==ids[0]
    a=s.get(p,'asset',ids[0]);a['sha256']='replaced';s.put(p,'asset',a,a['id'])
    assert FootageSearch(s,None).search(p,'桌上收纳盒',snap,sku_id='A')==[]


def test_selected_fragment_limits_actual_frames_and_text(tmp_path):
    import subprocess
    from creative_agent.media import MediaService
    from creative_agent.contracts import ModelBinding
    from creative_agent.video_understanding import VideoUnderstanding,plan_video_route
    s=Store(tmp_path/'db');p=s.create_project('p')['id'];media=MediaService(s,tmp_path/'media')
    path=tmp_path/'clip.mp4'
    subprocess.run([media.ffmpeg,'-y','-loglevel','error','-f','lavfi','-i','color=c=blue:s=48x48:r=4:d=2','-c:v','libx264',str(path)],check=True)
    a=media.ingest(p,path);a['metadata']['transcript_source']='user_import';a['metadata']['subtitles']=[{'start_ms':0,'end_ms':500,'text':'不能发送的前半段'},{'start_ms':1000,'end_ms':1500,'text':'选择的片段'}];s.put(p,'asset',a,a['id'])
    binding=ModelBinding(provider='openai',role='strategy',label='GPT',model_id='gpt-4.1')
    data=VideoUnderstanding(s,media).prepare(p,a['id'],binding,plan=plan_video_route(binding),interval={'start_ms':1000,'end_ms':1750})
    e=data['evidence']
    assert e['observed_timestamps_ms'] and all(1000<=t<1750 for t in e['observed_timestamps_ms'])
    assert [r['text'] for r in e['subtitles']]==['选择的片段']
    assert e['requested_range']=={'start_ms':1000,'end_ms':1750}


def test_output_count_is_not_confused_with_number_of_reference_videos():
    from creative_agent.creative_tasks import TaskSpec
    task=TaskSpec.infer('基于12条素材，写10条不同方向的30秒脚本')
    assert task.count==10 and task.target_duration_ms==30000


def test_corrected_factual_issue_can_be_resolved_without_losing_manual_notes(core):
    s,svc,p,q=core
    facts=svc.save_facts(p,0,{'product':'盒子','claims':[{'fact_id':'color','statement':'白色','source':'说明书','confirmed':True}]})
    body={**full(),'product_facts_snapshot':facts}
    body['rows'][0].update(dialogue='承重50公斤',claim_ids=['color'])
    draft=svc.create(p,body)
    assert draft['unresolved_items']
    fixed=deepcopy(draft['rows']);fixed[0].update(dialogue='这是白色的收纳盒',claim_ids=['color'])
    saved=svc.revise(p,draft['id'],1,{'title':draft['title'],'objective':draft['objective'],'rows':fixed,'confirmed':True})
    assert saved['confirmed'] and saved['unresolved_items']==[]


def test_typed_parse_cells_reject_wrong_values_and_keep_observation_basis(core):
    from creative_agent.analysis_templates import Templates
    s,svc,p,q=core
    t=Templates(s);template=t.save(p,{'title':'镜头表','fields':['物体数'],'definitions':[{'field_id':'count','name':'物体数','type':'number','level':'shot','basis':'observed','unit':'个'}]})
    assert t.cells(template['definitions'],{'物体数':2})['count']=={'value':2,'basis':'observed','state':'available','unit':'个'}
    with pytest.raises(ValueError):t.cells(template['definitions'],{'物体数':'很多'})


def test_parser_persists_typed_template_and_shot_cells(core,tmp_path):
    from creative_agent.workflows import Workflows
    from creative_agent.media import MediaService
    from PIL import Image
    s,svc,p,q=core;media=MediaService(s,tmp_path/'media');path=tmp_path/'still.png';Image.new('RGB',(16,16)).save(path)
    a=media.ingest(p,path)
    class Client:
        def chat(self,*args,**kwargs):return {'data':{'fields':{'开头':'展示包装'},'scenes':[]}}
    record=Workflows(s,media).parse(p,[a['id']],['开头'],Client())
    assert record['template_snapshot']['schema_version']=='2.1'
    assert list(record['rows'][0]['cells'].values())[0]['value']=='展示包装'


def test_commercial_packaging_requires_real_public_config_and_contains_no_secrets(tmp_path):
    import importlib.util
    path=Path(__file__).resolve().parents[1]/'scripts/package_config.py'
    assert path.is_file(),'packaging must validate commercial inputs before media preparation'
    spec=importlib.util.spec_from_file_location('package_config',path);mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    with pytest.raises(ValueError):mod.commercial_inputs(tmp_path,None)
    file=tmp_path/'public.json';file.write_text(json.dumps({'required':True,'public_key':'bad','service_url':'http://localhost'}))
    with pytest.raises(ValueError):mod.commercial_inputs(tmp_path,file)
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    from creative_agent.commercial_license import _encode
    file.write_text(json.dumps({'required':True,'public_key':_encode(Ed25519PrivateKey.generate().public_key().public_bytes_raw()),'service_url':'https://license.example.test'}))
    result=mod.commercial_inputs(tmp_path,file)
    assert (result/'commercial-license.json').is_file() and (result/'commercial-release.json').is_file()
    assert 'private' not in (result/'commercial-license.json').read_text()


def test_generated_script_cannot_omit_trusted_sku_to_reuse_wrong_item(core,tmp_path):
    from creative_agent.agent import AnalysisAgent
    s,svc,p,q=core
    facts=svc.save_facts(p,0,{'product':'A款','sku_id':'A'})
    f=tmp_path/'wrong.mp4';f.write_bytes(b'locator')
    a=s.put(p,'asset',{'source_kind':'own','path':str(f),'media_type':'video','metadata':{'duration_ms':5000,'sku_id':'B'}})
    script=full();script['rows'][0]['own_references']=[{'asset_id':a['id'],'start_ms':0,'end_ms':3000}]
    snap={'own_assets':[{'id':a['id']}]}
    with pytest.raises(ValueError,match='SKU|款式'):
        AnalysisAgent(s,None).prepare_scripts(p,{'result':{'scripts':[script]}},facts,snap,{'asset_ids':[]})


def test_subtitle_only_claims_need_the_same_fact_checks(core):
    s,svc,p,q=core
    body=full();body.update(confirmed=True);body['rows'][0].update(dialogue='',subtitle='到手价9.9元，承重50公斤')
    with pytest.raises(ValueError,match='事实|依据|核实'):svc.create(p,body)


def test_requested_duration_mismatch_is_not_complete():
    from creative_agent.creative_tasks import TaskSpec,validate_delivery
    result=validate_delivery({'scripts':[full()]},TaskSpec(output_type='scripts',target_duration_ms=30000))
    assert result['delivery_progress']['state']=='partial'
    assert result['duration_issues']


def test_selected_range_never_sends_outside_transcript_or_history_to_any_model(core,tmp_path):
    import subprocess
    from creative_agent.agent import AnalysisAgent
    from creative_agent.contracts import ModelBinding,AnalysisScope
    from creative_agent.media import MediaService
    from creative_agent.video_understanding import plan_video_route
    s,svc,p,q=core;media=MediaService(s,tmp_path/'media');path=tmp_path/'scope.mp4'
    subprocess.run([media.ffmpeg,'-y','-loglevel','error','-f','lavfi','-i','color=c=blue:s=48x48:r=4:d=2','-c:v','libx264',str(path)],check=True)
    a=media.ingest(p,path);a['metadata'].update(transcript_source='user_import',transcript='OUTSIDE_SENTINEL 全文',transcript_history=[{'text':'HISTORY_SENTINEL'}],subtitles=[{'start_ms':0,'end_ms':500,'text':'OUTSIDE_SENTINEL'},{'start_ms':1000,'end_ms':1500,'text':'选择的内容'}]);s.put(p,'asset',a,a['id'])
    class Model:
        binding=ModelBinding(provider='openai',role='strategy',label='GPT',model_id='gpt-4.1')
        def __init__(self):self.prompts=[]
        def chat(self,system,question,**kwargs):
            self.prompts.append(question)
            if question.startswith('逐条观察此素材'):
                import json
                supplied=json.loads(question.split('证据ID与帧时间：\n',1)[1].split('\n除',1)[0]);frame=supplied['input_manifest']['frames'][0]
                return {'data':{'summary':'抽样画面','observations':[{'claim_kind':'visual_observation','statement':'画面显示蓝色','interpretation':'observed','support':[{'frame_id':frame['frame_id'],'time_ms':frame['time_ms']}]}]},'usage':{}}
            return {'data':{'summary':'比较选中段落','findings':[],'result':{'comparisons':['展示选择的动作'],'recommendations':['比较顺序']}},'usage':{}}
    client=Model();scope=AnalysisScope(project_id=p,asset_ids=[a['id']],selected_segments=[{'asset_id':a['id'],'start_ms':1000,'end_ms':1750}])
    AnalysisAgent(s,media).analyze(scope,'directed','比较所选片段',client,video_route=plan_video_route(client.binding))
    assert len(client.prompts)==2
    assert all('OUTSIDE_SENTINEL' not in text and 'HISTORY_SENTINEL' not in text for text in client.prompts), [(i,text[max(0,text.find('OUTSIDE_SENTINEL')-80):text.find('OUTSIDE_SENTINEL')+120]) for i,text in enumerate(client.prompts) if 'OUTSIDE_SENTINEL' in text]


def test_script_without_existing_shot_or_reshoot_cannot_confirm(core):
    s,svc,p,q=core
    body=full();body['confirmed']=True;body['rows'][0]['reshoot']=''
    with pytest.raises(ValueError,match='补拍|素材'):svc.create(p,body)


def test_editing_confirmed_script_without_explicit_reconfirmation_creates_draft(core):
    s,svc,p,q=core
    saved=svc.create(p,{**full(),'confirmed':True})
    edited=svc.revise(p,saved['id'],1,{'rows':[{**saved['rows'][0],'dialogue':'人工修改后的台词'}]})
    assert edited['confirmed'] is False
    assert svc.version(p,saved['id'],1)['confirmed'] is True


def test_excel_keeps_formula_looking_text_literal(core):
    from openpyxl import load_workbook
    s,svc,p,q=core;body=full('=SUM(1,2)');body['rows'][0].update(dialogue='=1+1',notes='=1+2')
    script=svc.create(p,body);book=load_workbook(io.BytesIO(svc.export(p,script['id'],'xlsx')))
    assert book.active['C3'].value=='=1+1' and book.active['C3'].data_type=='s'
    assert book['制作说明与依据']['B2'].data_type=='s'


def test_selected_history_feedback_reaches_next_creation_and_freezes(core):
    from creative_agent.agent import AnalysisAgent
    from creative_agent.contracts import AnalysisScope
    from creative_agent.creative_tasks import TaskSpec
    s,svc,p,q=core;script=svc.create(p,{**full(),'usage_status':'rejected','feedback_reason':'不要夸张叫卖'})
    class Model:
        def __init__(self):self.prompts=[]
        def chat(self,system,question,**kwargs):
            self.prompts.append(json.loads(question))
            return {'data':{'summary':'新的表达','findings':[],'result':{'scripts':[full()]}},'usage':{}}
    client=Model();agent=AnalysisAgent(s,None)
    session=agent.analyze(AnalysisScope(project_id=p),'inspiration','写一个脚本',client,discussion_only=True,task_spec=TaskSpec(output_type='scripts',history_script_ids=[script['id']]))
    history=client.prompts[-1]['context']['creative_history']
    assert history[0]['feedback_reason']=='不要夸张叫卖'
    svc.revise(p,script['id'],1,{**full(),'feedback_reason':'后来改了'})
    agent.followup(p,session['id'],'继续修改',client)
    assert client.prompts[-1]['creative_history'][0]['feedback_reason']=='不要夸张叫卖'


def test_feature_revocation_stops_second_model_request_in_running_job(tmp_path,monkeypatch):
    import httpx
    from fastapi.testclient import TestClient
    from creative_agent.api import create_app
    from creative_agent.commercial_license import LicenseRuntime,sign_lease
    from creative_agent.license_service import LicenseDatabase,LicenseService
    from cryptography.hazmat.primitives.asymmetric.ed25519 import Ed25519PrivateKey
    from test_commercial_license import Vault
    from PIL import Image
    key=Ed25519PrivateKey.generate();clock=[1800000000.0];vault=Vault()
    runtime=LicenseRuntime(public_key=key.public_key(),service_url='https://license.test',cache_path=tmp_path/'lease.json',vault=vault,now=lambda:clock[0])
    svc=LicenseService(LicenseDatabase(tmp_path/'licenses'),key,now=lambda:clock[0]);account=svc.issue('test','strong-test-password',3)
    lease=svc.login(account['id'],'strong-test-password',runtime.device_id);runtime.accept(lease)
    monkeypatch.setattr(runtime,'start',lambda:None)
    app=create_app(tmp_path/'data',token='test',secret_store=vault,run_jobs=False,license_runtime=runtime)
    sent=[]
    def transport(request):
        sent.append(request.url.path)
        claims=runtime._claims();claims['features']=[f for f in claims['features'] if f!='creative.generate']
        runtime.accept({**lease,'lease':sign_lease(key,claims)})
        answer={'summary':'首条画面观察','findings':[]}
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps(answer)}}],'usage':{}})
    original=httpx.Client
    with TestClient(app,base_url='http://127.0.0.1:8765') as client:
        h={'Authorization':'Bearer test'};p=client.post('/api/projects',headers=h,json={'name':'test','budget':'10'}).json()['id']
        m=client.post(f'/api/projects/{p}/models',headers=h,json={'provider':'openai','role':'strategy','label':'test','model_id':'gpt-4.1','api_key':'test-not-a-real-key'}).json()
        f=tmp_path/'image.png';Image.new('RGB',(16,16)).save(f);a=app.state.media.ingest(p,f)
        job=client.post(f'/api/projects/{p}/actions',headers=h,json={'asset_ids':[a['id']],'question':'比较画面','quote':'1','mode':'directed','idempotency_key':'feature-check','binding_id':m['id']}).json()
        monkeypatch.setattr('creative_agent.providers.httpx.Client',lambda **kwargs:original(transport=httpx.MockTransport(transport)))
        app.state.execute(p,job['id'])
        assert sent==['/chat/completions'] or sent==['/v1/chat/completions']
        assert app.state.store.job(p,job['id'])['state']=='needs_attention'


def test_large_script_request_batches_and_keeps_completed_round_on_failure(core):
    from creative_agent.agent import AnalysisAgent
    from creative_agent.contracts import AnalysisScope
    from creative_agent.creative_tasks import TaskSpec
    s,svc,p,q=core
    class Model:
        def __init__(self,fail=False):self.n=0;self.fail=fail
        def chat(self,*args,**kwargs):
            self.n+=1
            if self.fail and self.n==2:raise ValueError('外部请求结果不明，不能重发')
            return {'data':{'summary':'两条完整脚本','findings':[],'result':{'scripts':[full(f'方向{self.n}A'),full(f'方向{self.n}B')]}},'usage':{}}
    spec=TaskSpec(output_type='scripts',count=4,max_generation_rounds=2)
    client=Model();agent=AnalysisAgent(s,None)
    result=agent.analyze(AnalysisScope(project_id=p),'inspiration','写四条脚本',client,discussion_only=True,task_spec=spec)
    assert client.n==2 and len(result['turns'][0]['answer']['result']['scripts'])==4
    assert result['turns'][0]['answer']['result']['delivery_progress']['state']=='complete'
    client=Model(fail=True)
    with pytest.raises(ValueError,match='结果不明'):agent.analyze(AnalysisScope(project_id=p),'inspiration','再写四条',client,discussion_only=True,task_spec=spec)
    partial=[x for x in s.list(p,'session') if x['turns'][0]['question']=='再写四条'][0]
    assert len(partial['turns'][0]['answer']['result']['scripts'])==2
    assert partial['turns'][0]['answer']['result']['delivery_progress']['state']=='partial'


def test_followup_batches_requested_scripts_and_checkpoints_completed_work(core):
    from creative_agent.agent import AnalysisAgent
    from creative_agent.contracts import AnalysisScope
    s,svc,p,q=core
    class Model:
        def __init__(self):self.n=0
        def chat(self,*args,**kwargs):
            self.n+=1
            return {'data':{'summary':'完整脚本','findings':[],'result':{'scripts':[full(f'新方向{self.n}A'),full(f'新方向{self.n}B')]}},'usage':{}}
    agent=AnalysisAgent(s,None);client=Model()
    session=agent.analyze(AnalysisScope(project_id=p),'inspiration','写两个脚本',client,discussion_only=True)
    client.n=0
    result=agent.followup(p,session['id'],'再写四个完整脚本',client)
    assert client.n==2
    assert result['turns'][-1]['answer']['result']['delivery_progress']['completed']==4


def test_scene_index_request_limit_counts_scenes_not_files(tmp_path,monkeypatch):
    import httpx
    from fastapi.testclient import TestClient
    from creative_agent.api import create_app
    from test_commercial_license import Vault
    app=create_app(tmp_path/'data',token='test',secret_store=Vault(),run_jobs=False)
    original=httpx.Client;sent=[]
    def transport(request):
        sent.append(request.url.path)
        return httpx.Response(200,json={'data':[{'embedding':[1.0,0.0]}],'usage':{}})
    with TestClient(app,base_url='http://127.0.0.1:8765') as c:
        h={'Authorization':'Bearer test'};p=c.post('/api/projects',headers=h,json={'name':'test','budget':'10'}).json()['id']
        b=c.post(f'/api/projects/{p}/models',headers=h,json={'provider':'openai','role':'embedding','label':'test','model_id':'text-embedding-3-small','api_key':'test-not-a-real-key'}).json()
        s=app.state.store
        a=s.put(p,'asset',{'title':'三镜素材','media_type':'video','inspection':'ready','metadata':{'duration_ms':3000,'transcript_source':'user_import','subtitles':[{'start_ms':0,'end_ms':3000,'text':'整理桌面'}]}})
        for i in range(3):s.put(p,'scene',{'asset_id':a['id'],'start_ms':i*1000,'end_ms':(i+1)*1000})
        job=c.post(f'/api/projects/{p}/actions',headers=h,json={'kind':'scene_embedding','asset_ids':[a['id']],'binding_id':b['id'],'quote':'1','idempotency_key':'scenes'}).json()
        monkeypatch.setattr('creative_agent.providers.httpx.Client',lambda **kwargs:original(transport=httpx.MockTransport(transport)))
        app.state.execute(p,job['id'])
        assert len(sent)==3
        assert s.job(p,job['id'])['state']=='succeeded'


@pytest.mark.parametrize('timeout_after_cancel',[False,True])
def test_queue_freezes_facts_and_cancel_keeps_first_batch(tmp_path,monkeypatch,timeout_after_cancel):
    import httpx
    from fastapi.testclient import TestClient
    from creative_agent.api import create_app
    from test_commercial_license import Vault
    app=create_app(tmp_path/'data',token='test',secret_store=Vault(),run_jobs=False)
    original=httpx.Client;sent=[];job={}
    def transport(request):
        data=json.loads(json.loads(request.content)['messages'][-1]['content'][0]['text']);sent.append(data)
        if not timeout_after_cancel or len(sent)==2:app.state.store.cancel(p,job['id'])
        if timeout_after_cancel and len(sent)==2:raise httpx.ReadTimeout('controlled timeout')
        answer={'summary':'两条脚本已完成','findings':[],'result':{'scripts':[full('第一条'),full('第二条')]}}
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps(answer)}}],'usage':{}})
    with TestClient(app,base_url='http://127.0.0.1:8765') as c:
        h={'Authorization':'Bearer test'};p=c.post('/api/projects',headers=h,json={'name':'test','budget':'10'}).json()['id']
        b=c.post(f'/api/projects/{p}/models',headers=h,json={'provider':'openai','role':'strategy','label':'test','model_id':'gpt-4.1','api_key':'test-not-a-real-key'}).json()
        scripts=Scripts(app.state.store)
        scripts.save_facts(p,0,{'product':'甲商品','sku_id':'A'})
        job.update(c.post(f'/api/projects/{p}/actions',headers=h,json={'kind':'analysis','discussion_only':True,'question':'写四个完整脚本','binding_id':b['id'],'quote':'1','idempotency_key':'frozen-cancel'}).json())
        scripts.save_facts(p,1,{'product':'乙商品','sku_id':'B'})
        monkeypatch.setattr('creative_agent.providers.httpx.Client',lambda **kwargs:original(transport=httpx.MockTransport(transport)))
        app.state.execute(p,job['id'])
        assert sent[0]['context']['product_facts']['sku_id']=='A'
        saved=app.state.store.job(p,job['id'])
        assert saved['state']=='cancelled' and saved['actual_cost'] is None
        session=app.state.store.get(p,'session',saved['result']['record_id'])
        assert len(session['turns'][0]['answer']['result']['scripts'])==2
        assert len(sent)==(2 if timeout_after_cancel else 1)


def test_task_quantity_and_duration_follow_requested_output_not_reference_length():
    from creative_agent.creative_tasks import TaskSpec
    spec=TaskSpec.infer('参考这些素材前3秒的开头，写十一个30秒的完整脚本')
    assert spec.count==11 and spec.target_duration_ms==30000
    spec=TaskSpec.infer('写十三条脚本，每条20秒')
    assert spec.count==13 and spec.target_duration_ms==20000
