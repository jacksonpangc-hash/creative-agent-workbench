"""Real, version-specific shooting documents. No HTML renamed as Word."""
import io
import json
from pathlib import Path
import threading

_FONT_LOCK=threading.Lock()
from html import escape
from zipfile import ZipFile, ZIP_DEFLATED

COLUMNS=['镜头','时间（秒）','台词','字幕','画面','机位','语气','道具','参考表达','自有镜头','补拍','商品依据']


def rows(script):
    # Labels belong to this saved conclusion, never the current asset registry.
    titles={e.get('asset_id') or e.get('id'):e['title']
            for e in (script.get('source') or {}).get('evidence',[])
            if isinstance(e.get('title'),str) and e['title'].strip()}
    def refs(values):
        lines=[]
        modes={'normal':'正常速度','freeze_frame':'定格','slow_motion':'慢放','loop':'循环'}
        for r in values:
            aid=r['asset_id']
            label=f"{titles[aid]}（{aid}）" if aid in titles else aid
            line=f"{label} · {r.get('start_ms',0)}–{r.get('end_ms','整图')} ms"
            if r.get('usage_mode') and (r['usage_mode']!='normal' or r.get('use_duration_ms')):
                line+='；展示：'+modes[r['usage_mode']]+' '+str(r['use_duration_ms'])+' ms'
                if r.get('usage_note'):line+='；'+r['usage_note']
                if r['usage_mode']!='normal':line+='；原始观察范围不变'
            lines.append(line)
        return '\n'.join(lines)
    def reshoot(row):
        text=row.get('reshoot','')
        if row.get('reshoot_duration_ms') is not None:text+='；计划补拍 '+str(row['reshoot_duration_ms'])+' ms'
        return text
    return [[str(i+1),f"{r['start_ms']/1000:g}–{r['end_ms']/1000:g}",r['dialogue'],r.get('subtitle',''),r['visual'],r.get('camera',''),r.get('tone',''),
             '、'.join(r.get('props',[])),refs(r.get('references',[])),refs(r.get('own_references',[])),reshoot(r),
             '、'.join(r.get('claim_ids',[]))] for i,r in enumerate(script['rows'])]



def evidence_lines(script):
    facts=script.get('product_facts_snapshot',{})
    lines=[]
    if facts:
        lines.append('商品：'+str(facts.get('product',''))+' · SKU：'+str(facts.get('sku_id',''))+' · 事实版本：'+str(facts.get('version',0)))
        for claim in facts.get('claims',[]):
            lines.append(str(claim.get('fact_id',''))+'：'+str(claim.get('statement',''))+'；依据：'+str(claim.get('source','') or '未提供')+'；'+('已确认' if claim.get('confirmed') else '待核实')+('；有效至'+str(claim['valid_until']) if claim.get('valid_until') else ''))
        if facts.get('price_and_promotion'):lines.append('价格与促销：'+str(facts['price_and_promotion'])+'；有效期：'+str(facts.get('promotion_valid_until') or '未提供'))
    for i,row in enumerate(script['rows']):
        for claim in row.get('factual_claims',[]):lines.append(f'镜头{i+1}说法：'+str(claim.get('statement',''))+'；依据ID：'+'、'.join(claim.get('fact_ids',[]))+'；'+('人工已核对' if claim.get('reviewed') else '支持关系待核实'))
    return lines


def markdown(script):
    text=f"# {script['title']}\n\n版本：{script.get('version',1)} · {'已确认' if script.get('confirmed') else '草稿'}\n\n{script['objective']}\n\n"
    cell=lambda x:str(x).replace('|','\\|').replace('\n','；')
    text+='| '+' | '.join(COLUMNS)+' |\n| '+' | '.join(['---']*len(COLUMNS))+' |\n'
    text+='\n'.join('| '+' | '.join(cell(x) for x in r)+' |' for r in rows(script))
    if script.get('unresolved_items'):text+='\n\n## 待核实\n\n'+'\n'.join('- '+x for x in script['unresolved_items'])
    if script.get('production_notes'):text+='\n\n## 制作备注\n\n'+'\n'.join('- '+x for x in script['production_notes'])
    for i,r in enumerate(script['rows']):
        if r.get('notes'):text+=f"\n\n镜头 {i+1} 制作说明：{r['notes']}"
    if evidence_lines(script):text+='\n\n## 商品依据与核对记录\n\n'+'\n'.join('- '+x for x in evidence_lines(script))
    return text


def export_script(script,format):
    if format=='json':return json.dumps(script,ensure_ascii=False,indent=2).encode()
    if format=='md':return markdown(script).encode()
    output=io.BytesIO()
    if format=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        book=Workbook();sheet=book.active;sheet.title='拍摄脚本'
        sheet.append([script['title'],f"v{script.get('version',1)}",script['objective']]);sheet.append(COLUMNS)
        for row in rows(script):sheet.append(row)
        for row in sheet:
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        for cell in sheet[2]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='6744D4')
        for i in range(1,len(COLUMNS)+1):sheet.column_dimensions[get_column_letter(i)].width=32 if i in (3,5,9,10,11) else 18
        sheet.freeze_panes='A3';sheet.auto_filter.ref=f'A2:L{sheet.max_row}'
        detail=book.create_sheet('制作说明与依据');detail.append(['镜头','制作说明','商品说法'])
        for i,r in enumerate(script['rows']):detail.append([i+1,r.get('notes',''),json.dumps(r.get('factual_claims',[]),ensure_ascii=False)])
        detail.append(['待核实','；'.join(script.get('unresolved_items',[]))])
        for note in script.get('production_notes',[]):detail.append(['制作备注',note])
        for line in evidence_lines(script):detail.append(['商品依据',line])
        detail.column_dimensions['A'].width=18;detail.column_dimensions['B'].width=72;detail.column_dimensions['C'].width=60
        for row in detail:
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        for page in book:
            for row in page:
                for cell in row:
                    if isinstance(cell.value,str):cell.data_type='s'
        book.save(output)
    elif format=='docx':
        def para(text,bold=False,title=False):
            pstyle='<w:pStyle w:val="Title"/>' if title else ''
            return '<w:p><w:pPr>'+pstyle+'<w:spacing w:after="120" w:line="300" w:lineRule="auto"/>'+('<w:keepNext/>' if bold else '')+'</w:pPr><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:eastAsia="Microsoft YaHei"/><w:color w:val="000000"/><w:sz w:val="'+('36' if title else '22')+'"/>'+('<w:b/>' if bold else '')+'</w:rPr><w:t xml:space="preserve">'+escape(str(text))+'</w:t></w:r></w:p>'
        body=para(script['title'],True,True)+para(f"版本 {script.get('version',1)} · "+('已确认' if script.get('confirmed') else '草稿'))+para(script['objective'])
        # A per-shot vertical layout remains readable in Word instead of a 12-column squeezed table.
        for i,r in enumerate(rows(script)):
            body+=para('镜头 '+str(i+1)+' · '+r[1]+' 秒',True)
            for k,value in zip(COLUMNS[2:],r[2:]):
                if value:body+=para(k+'：'+value)
            if script['rows'][i].get('notes'):body+=para('制作说明：'+script['rows'][i]['notes'])
        if script.get('unresolved_items'):body+=para('待核实',True)+''.join(para(x) for x in script['unresolved_items'])
        if script.get('production_notes'):body+=para('制作备注',True)+''.join(para(x) for x in script['production_notes'])
        if evidence_lines(script):body+=para('商品依据与核对记录',True)+''.join(para(x) for x in evidence_lines(script))
        xml='<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'+body+'<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/></w:sectPr></w:body></w:document>'
        with ZipFile(output,'w',ZIP_DEFLATED) as z:
            z.writestr('[Content_Types].xml','<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            z.writestr('_rels/.rels','<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
            z.writestr('word/document.xml',xml)
    elif format=='pdf':
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,KeepTogether
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        with _FONT_LOCK:
            if 'CreativeSansSC' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('CreativeSansSC',str(Path(__file__).with_name('fonts')/'CreativeSansSC-Regular.ttf')))
        style=ParagraphStyle('Chinese',fontName='CreativeSansSC',fontSize=11,leading=18,spaceAfter=6,wordWrap='CJK')
        title=ParagraphStyle('Title',parent=style,fontSize=18,leading=25,spaceAfter=15)
        parts=[Paragraph(escape(script['title']),title),Paragraph(f"版本 {script.get('version',1)}",style),Paragraph(escape(script['objective']),style)]
        for i,r in enumerate(rows(script)):
            parts.extend([Spacer(1,12),Paragraph('镜头 '+str(i+1)+' · '+r[1]+' 秒',style)])
            for k,value in zip(COLUMNS[2:],r[2:]):
                if value:parts.append(Paragraph(escape(k+'：'+value).replace('\n','<br/>'),style))
            if script['rows'][i].get('notes'):parts.append(Paragraph(escape('制作说明：'+script['rows'][i]['notes']),style))
        if script.get('unresolved_items'):parts.append(Paragraph(escape('待核实：'+'；'.join(script['unresolved_items'])),style))
        if script.get('production_notes'):parts.extend([Paragraph('制作备注',style),*[Paragraph(escape(x),style) for x in script['production_notes']]])
        if evidence_lines(script):parts.extend([Spacer(1,12),Paragraph('商品依据与核对记录',style),*[Paragraph(escape(x),style) for x in evidence_lines(script)]])
        SimpleDocTemplate(output,title=script['title'],leftMargin=48,rightMargin=48,topMargin=48,bottomMargin=48).build(parts)
    else:raise ValueError('不支持的脚本导出格式')
    return output.getvalue()
